"""PDF and DOCX ingestion.

These paths are what a real student hits - the .txt fixtures never exercise
pypdf or python-docx, so without these tests the most likely real-world entry
point is the least tested one.
"""

from __future__ import annotations

import io

import pytest

from s2s.extract import heuristic
from s2s.ingest import loader, structure
from s2s.ingest.loader import IngestError
from tests.pdf_builder import build_pdf

SYLLABUS_LINES = [
    "Paper BCH 3.2: Financial Management",
    "Unit I: Working Capital Management",
    "Operating cycle and cash conversion cycle. Estimation of working capital",
    "requirements. Ratio analysis for liquidity: current ratio, quick ratio.",
    "Unit II: Capital Budgeting",
    "Net Present Value (NPV), Internal Rate of Return (IRR) and payback period.",
]


# ----------------------------------------------------------------- PDF -----

@pytest.fixture
def syllabus_pdf() -> bytes:
    return build_pdf(SYLLABUS_LINES)


def test_pdf_text_is_extracted(syllabus_pdf):
    syllabus = loader.from_pdf(syllabus_pdf, "generated.pdf")
    assert "Working Capital Management" in syllabus.text
    assert "Internal Rate of Return" in syllabus.text


def test_pdf_parses_into_units(syllabus_pdf):
    syllabus = loader.from_pdf(syllabus_pdf, "generated.pdf")
    tree = structure.parse(syllabus.text)
    assert tree.unit_count == 2


def test_pdf_yields_competencies_with_verbatim_spans(syllabus_pdf):
    syllabus = loader.from_pdf(syllabus_pdf, "generated.pdf")
    tree = structure.parse(syllabus.text)
    competencies = heuristic.extract(syllabus, tree, semester=3)
    assert competencies
    for competency in competencies:
        assert competency.source_span in syllabus.text


def test_pdf_dispatches_through_from_bytes(syllabus_pdf):
    syllabus = loader.from_bytes(syllabus_pdf, "my syllabus.PDF")
    assert "Capital Budgeting" in syllabus.text


def test_scanned_pdf_raises_an_actionable_error():
    """An image-only PDF must say so, not silently return nothing."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(IngestError, match="(?i)scan"):
        loader.from_pdf(buffer.getvalue())


# ---------------------------------------------------------------- DOCX -----

def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Paper SOC 202: Methods of Sociological Enquiry")
    document.add_paragraph("Unit I: Qualitative Research Methods")
    document.add_paragraph(
        "Semi-structured interviewing, thematic analysis, open coding and "
        "building a qualitative codebook from field data."
    )
    document.add_paragraph("Unit II: Survey Design")
    document.add_paragraph(
        "Questionnaire construction, Likert scales, sampling and piloting."
    )
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Credits"
    table.rows[0].cells[1].text = "6"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_text_and_tables_are_extracted():
    syllabus = loader.from_docx(_docx_bytes(), "x.docx")
    assert "Qualitative Research Methods" in syllabus.text
    assert "Credits" in syllabus.text, "table cells should be read too"


def test_docx_parses_into_units_and_competencies():
    syllabus = loader.from_docx(_docx_bytes(), "x.docx")
    tree = structure.parse(syllabus.text)
    assert tree.unit_count == 2

    competencies = heuristic.extract(syllabus, tree, semester=2)
    assert competencies
    for competency in competencies:
        assert competency.source_span in syllabus.text


def test_docx_dispatches_through_from_bytes():
    syllabus = loader.from_bytes(_docx_bytes(), "Syllabus.DOCX")
    assert "Survey Design" in syllabus.text


# ------------------------------------------------------------ dispatch -----

def test_unsupported_extension_is_rejected():
    with pytest.raises(IngestError, match="(?i)unsupported"):
        loader.from_bytes(b"data", "deck.pptx")


def test_utf16_text_is_decoded():
    """The requirements.txt lesson: encoding detection must not assume UTF-8."""
    body = (
        "Unit I: Working Capital Management\n"
        "Operating cycle and cash conversion cycle analysis for firms. " * 3
    )
    syllabus = loader.from_bytes(body.encode("utf-16"), "syllabus.txt")
    assert "Working Capital" in syllabus.text
