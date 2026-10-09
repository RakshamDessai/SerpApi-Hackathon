"""Stage 0a: get raw text out of whatever the student uploaded.

Supports PDF, DOCX, TXT and direct paste. Scanned-image PDFs are out of scope;
we detect "no extractable text" and say so rather than silently returning an
empty string.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

MIN_USEFUL_CHARS = 120


class IngestError(Exception):
    """Raised when a document yields no usable text."""


@dataclass
class RawSyllabus:
    text: str
    source: str
    tree: "CurriculumNode | None" = field(default=None)

    @property
    def char_count(self) -> int:
        return len(self.text)


def _normalise(text: str) -> str:
    """Collapse the whitespace noise typical of PDF extraction.

    Offsets are computed against this normalised text, so Competency.span_offset
    stays valid for UI highlighting.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ")          # non-breaking space
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def from_text(raw: str, source: str = "pasted text") -> RawSyllabus:
    text = _normalise(raw)
    if len(text) < MIN_USEFUL_CHARS:
        raise IngestError(
            f"Only {len(text)} characters of text found - too short to parse as a "
            "syllabus. Paste at least a unit listing."
        )
    return RawSyllabus(text=text, source=source)


def from_pdf(data: bytes, source: str = "uploaded PDF") -> RawSyllabus:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:                       # a damaged page should not kill the run
            pages.append("")
    text = _normalise("\n".join(pages))

    if len(text) < MIN_USEFUL_CHARS:
        raise IngestError(
            "No extractable text in that PDF. It is most likely a scan - S2S does "
            "not run OCR. Paste the unit list instead, or use a text-based PDF."
        )
    return RawSyllabus(text=text, source=source)


def from_docx(data: bytes, source: str = "uploaded DOCX") -> RawSyllabus:
    import docx

    document = docx.Document(io.BytesIO(data))
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:                # unit lists are often tables
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = _normalise("\n".join(chunks))
    if len(text) < MIN_USEFUL_CHARS:
        raise IngestError("That DOCX contained no readable paragraphs or tables.")
    return RawSyllabus(text=text, source=source)


def from_bytes(data: bytes, filename: str) -> RawSyllabus:
    """Dispatch on file extension."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return from_pdf(data, source=filename)
    if suffix in {".docx", ".doc"}:
        return from_docx(data, source=filename)
    if suffix in {".txt", ".md", ".text", ""}:
        return from_text(_decode(data), source=filename)
    raise IngestError(f"Unsupported file type '{suffix}'. Use PDF, DOCX or TXT.")


def _decode(data: bytes) -> str:
    """Decode bytes of unknown encoding (the requirements.txt lesson)."""
    try:
        from charset_normalizer import from_bytes as cn_from_bytes

        best = cn_from_bytes(data).best()
        if best is not None:
            return str(best)
    except Exception:
        pass
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def from_fixture(path: Path) -> RawSyllabus:
    return from_text(path.read_text(encoding="utf-8"), source=path.stem)


# Imported late to avoid a circular import at module load.
from s2s.ingest.structure import CurriculumNode  # noqa: E402  (type only)
